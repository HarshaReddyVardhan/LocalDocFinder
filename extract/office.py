"""DOCX (chunked by heading) and PPTX (one chunk per slide), with embedded-image OCR."""
import io
import zipfile
from typing import List

import xxhash
from PIL import Image

import indexer_config as cfg
from . import Chunk, ExtractError, split_by_lines


def _image_chunk(data: bytes, page: int, seen: set, label: str):
    key = xxhash.xxh3_64_hexdigest(data)
    if key in seen:
        return None
    seen.add(key)
    try:
        with Image.open(io.BytesIO(data)) as im:
            im.load()
            if min(im.size) < cfg.MIN_IMAGE_PIXELS:
                return None
            from .image import describe
            body = describe(im, label)
    except Exception:
        return None
    return Chunk(body, "image", label, page=page) if body else None


def extract_docx(path: str) -> List[Chunk]:
    import docx
    try:
        document = docx.Document(path)
    except Exception as e:
        raise ExtractError(f"cannot open docx: {e}")

    chunks: List[Chunk] = []
    stack: List[tuple] = []
    cur_path, buf = "", []

    def flush():
        text = "\n".join(buf).strip()
        buf.clear()
        if not text:
            return
        if len(text) <= cfg.MAX_CHUNK_CHARS:
            chunks.append(Chunk(text, "doc", cur_path))
        else:
            for t, _, _ in split_by_lines(text.splitlines(), 1, max_chars=cfg.TARGET_CHUNK_CHARS * 2):
                chunks.append(Chunk(t, "doc", cur_path))

    for p in document.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith("heading") or style == "Title":
            flush()
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, txt))
            cur_path = " > ".join(t for _, t in stack)
            buf.append(txt)
        else:
            buf.append(txt)
    flush()

    for ti, table in enumerate(document.tables, 1):
        rows = []
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip()
                if t and (not cells or cells[-1] != t):
                    cells.append(t)
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            for t, _, _ in split_by_lines(rows, 1, max_chars=cfg.TARGET_CHUNK_CHARS * 2):
                chunks.append(Chunk(t, "doc", f"table {ti}"))

    # Embedded images (word/media/*)
    budget, seen = cfg.MAX_IMAGES_PER_DOC, set()
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if budget <= 0:
                    break
                if name.startswith("word/media/") and name.lower().endswith(
                        (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif")):
                    budget -= 1
                    c = _image_chunk(z.read(name), 0, seen, "embedded image")
                    if c:
                        chunks.append(c)
    except zipfile.BadZipFile:
        pass
    return chunks


def extract_pptx(path: str) -> List[Chunk]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        prs = Presentation(path)
    except Exception as e:
        raise ExtractError(f"cannot open pptx: {e}")
    chunks: List[Chunk] = []
    budget, seen = cfg.MAX_IMAGES_PER_DOC, set()

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    for n, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in walk(slide.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE and budget > 0:
                budget -= 1
                c = _image_chunk(sh.image.blob, n, seen, f"image on slide {n}")
                if c:
                    chunks.append(c)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append("Notes: " + notes)
        if texts:
            chunks.append(Chunk("\n".join(texts)[: cfg.MAX_CHUNK_CHARS], "doc", f"slide {n}", page=n))
    return chunks


def extract_office(path: str) -> List[Chunk]:
    return extract_docx(path) if path.lower().endswith(".docx") else extract_pptx(path)
