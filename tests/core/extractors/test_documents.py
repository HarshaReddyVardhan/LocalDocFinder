import io
from collections.abc import Callable
from pathlib import Path

import pymupdf
import pytest
from tests.core.extractors.conftest import FakeOcr, Writer, png_bytes

from localdoc_finder.core.extractors import pdf as pdf_mod
from localdoc_finder.core.extractors.base import ExtractError, ExtractorSet
from localdoc_finder.core.settings import ChunkingSettings, ImageSettings

Build = Callable[..., ExtractorSet]


def make_pdf(path: Path, pages: list[str], images: dict[int, bytes] | None = None) -> Path:
    doc = pymupdf.open()
    for number, text in enumerate(pages, 1):
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
        if images and number in images:
            page.insert_image(pymupdf.Rect(72, 200, 372, 500), stream=images[number])
    doc.save(path)
    doc.close()
    return path


class TestPdf:
    def test_text_per_page_with_page_numbers(
        self, extractors: ExtractorSet, tmp_path: Path
    ) -> None:
        pdf = make_pdf(tmp_path / "a.pdf", ["first page text " * 5, "second page text " * 5])
        chunks = extractors.extract(pdf)
        assert [(c.page, c.symbol) for c in chunks] == [(1, "p.1"), (2, "p.2")]
        assert "second page" in chunks[1].text

    def test_long_page_is_split(
        self, with_chunking: Callable[..., ExtractorSet], tmp_path: Path
    ) -> None:
        doc = pymupdf.open()
        page = doc.new_page()
        for i in range(70):
            page.insert_text((20, 20 + 11 * i), f"line {i} of a very long page", fontsize=8)
        doc.save(tmp_path / "long.pdf")
        doc.close()
        cfg = ChunkingSettings(max_chunk_chars=200, target_chunk_chars=100)
        chunks = with_chunking(cfg).extract(tmp_path / "long.pdf")
        assert len(chunks) > 1
        assert {c.page for c in chunks} == {1}

    def test_figure_text_is_ocrd_with_page(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        pdf = make_pdf(tmp_path / "d.pdf", ["intro text " * 5], images={1: png_bytes((300, 300))})
        chunks = build_with_ocr().extract(pdf)
        figure = next(c for c in chunks if c.kind == "image")
        assert figure.page == 1
        assert "OAuth login flow" in figure.text
        assert figure.symbol == "figure on p.1"

    def test_scanned_page_is_ocrd(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        ocr.text = "scanned contract clause seven"
        pdf = make_pdf(tmp_path / "scan.pdf", [""], images={1: png_bytes((600, 800))})
        chunks = build_with_ocr().extract(pdf)
        assert [(c.kind, c.page) for c in chunks] == [("doc", 1)]
        assert "clause seven" in chunks[0].text

    def test_a_scan_drawn_as_an_inline_image_is_ocrd(
        self,
        build_with_ocr: Build,
        tmp_path: Path,
        ocr: FakeOcr,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        ocr.text = "inline scanned page"
        pdf = make_pdf(tmp_path / "inline.pdf", [""], images={1: png_bytes((600, 800))})
        # An inline image is drawn by the page itself and missing from get_images().
        monkeypatch.setattr(pymupdf.Page, "get_images", lambda self, full=False: [])
        chunks = build_with_ocr().extract(pdf)
        assert [(c.kind, c.page) for c in chunks] == [("doc", 1)]
        assert "inline scanned page" in chunks[0].text

    def test_a_long_scan_is_read_past_the_picture_allowance(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        colours = ["red", "green", "blue", "gray", "black"]
        ocr.pages = [f"clause {n} of the {c} contract" for n, c in enumerate(colours, 1)]
        pdf = make_pdf(
            tmp_path / "long-scan.pdf",
            [""] * 5,
            images={n: png_bytes((600, 800), c) for n, c in enumerate(colours, 1)},
        )
        chunks = build_with_ocr(ImageSettings(max_per_doc=2)).extract(pdf)
        assert ocr.calls == 5
        assert [c.page for c in chunks] == [1, 2, 3, 4, 5]

    def test_ocr_budget_limits_work(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        pdf = make_pdf(
            tmp_path / "many.pdf",
            [""] * 5,
            images={
                n: png_bytes((600, 800), c)
                for n, c in enumerate(["red", "green", "blue", "gray", "black"], 1)
            },
        )
        build_with_ocr(ImageSettings(max_scanned_pages=2)).extract(pdf)
        assert ocr.calls == 2

    def test_scanned_pages_do_not_use_up_the_allowance_for_figures(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        doc = pymupdf.open()
        for color in ("red", "green"):  # two scanned pages
            page = doc.new_page()
            page.insert_image(pymupdf.Rect(0, 0, 400, 600), stream=png_bytes((600, 800), color))
        text_page = doc.new_page()  # then a normal page with a figure
        text_page.insert_text((72, 72), "regular page text " * 5)
        text_page.insert_image(
            pymupdf.Rect(72, 200, 372, 500), stream=png_bytes((300, 300), "blue")
        )
        target = tmp_path / "mixed.pdf"
        doc.save(target)
        doc.close()
        chunks = build_with_ocr(ImageSettings(max_per_doc=2, max_scanned_pages=2)).extract(target)
        assert any(c.kind == "image" and c.page == 3 for c in chunks)  # the figure was still read
        assert ocr.calls == 3

    def test_an_image_bomb_in_a_pdf_skips_that_image_not_the_file(
        self, build_with_ocr: Build, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from PIL import Image

        from localdoc_finder.core.extractors import pdf as pdf_module

        def bomb(*_args: object, **_kwargs: object) -> None:
            raise Image.DecompressionBombError("too many pixels")

        monkeypatch.setattr(pdf_module, "describe", bomb)
        pdf = make_pdf(
            tmp_path / "bomb.pdf", ["body text " * 10], images={1: png_bytes((300, 300))}
        )
        chunks = build_with_ocr().extract(pdf)
        assert [c.kind for c in chunks] == ["doc"]  # the text survives

    def test_duplicate_figures_are_ocrd_once(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        same = png_bytes()
        pdf = make_pdf(
            tmp_path / "dup.pdf", ["text " * 10, "more " * 10], images={1: same, 2: same}
        )
        build_with_ocr().extract(pdf)
        assert ocr.calls == 1

    def test_small_figures_are_skipped(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        pdf = make_pdf(tmp_path / "logo.pdf", ["text " * 10], images={1: png_bytes((40, 40))})
        build_with_ocr().extract(pdf)
        assert ocr.calls == 0

    def scans(self, tmp_path: Path, count: int, name: str = "scan.pdf") -> Path:
        colours = ["red", "green", "blue", "gray", "black"]
        return make_pdf(
            tmp_path / name,
            [""] * count,
            images={n: png_bytes((600, 800), colours[n - 1]) for n in range(1, count + 1)},
        )

    def test_a_watermark_on_every_page_is_stripped(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        ocr.pages = [
            "RESIDENTIAL LEASE AGREEMENT between the landlord and the tenant\n"
            "Scanned with CamScanner\nPage 1 of 2",
            "The tenant pays a monthly rent of 1450 dollars on the first day\n"
            "Scanned with CamScanner\nPage 2 of 2",
        ]
        chunks = build_with_ocr().extract(self.scans(tmp_path, 2))
        assert [c.page for c in chunks] == [1, 2]
        assert all("CamScanner" not in c.text and "Page" not in c.text for c in chunks)
        assert not any(c.low_content for c in chunks)

    def test_a_junk_scan_keeps_only_a_low_content_name_chunk(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        ocr.text = "Scanned with CamScanner"  # every page: nothing but the watermark
        chunks = build_with_ocr().extract(self.scans(tmp_path, 3, "scan_0042.pdf"))
        assert len(chunks) == 1
        assert chunks[0].low_content
        assert chunks[0].text == "PDF file: scan_0042.pdf (3 pages)"

    @pytest.mark.parametrize(
        ("text", "low"),
        [
            ("' . , i l | ~ - ii ;: .. ' ' Il 1 l. , lll iiii | ~~", True),  # dust read as text
            ("Scanned with CamScanner", True),  # too little to say anything
            ("This lease is made between the landlord and the tenant for twelve months", False),
        ],
    )
    def test_a_scanned_page_with_little_real_text_is_low_content(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr, text: str, low: bool
    ) -> None:
        ocr.text = text
        chunks = build_with_ocr().extract(self.scans(tmp_path, 1))
        assert [c.low_content for c in chunks] == [low]

    def test_scanned_pages_past_the_ocr_allowance_are_low_content(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        ocr.text = "This lease is made between the landlord and the tenant for twelve months"
        # Each page carries a scrap of real text (a stamp), too short to skip OCR.
        colours = ["red", "green", "blue"]
        pdf = make_pdf(
            tmp_path / "stamped.pdf",
            ["Ref A1", "Note B2", "Memo C3"],
            images={n: png_bytes((600, 800), c) for n, c in enumerate(colours, 1)},
        )
        chunks = build_with_ocr(ImageSettings(max_scanned_pages=1)).extract(pdf)
        assert [(c.page, c.low_content) for c in chunks] == [(1, False), (2, True), (3, True)]
        assert "Note B2" in chunks[1].text  # still there for keyword search

    def test_title_and_subject_go_into_the_first_page(
        self, extractors: ExtractorSet, tmp_path: Path
    ) -> None:
        pdf = make_pdf(tmp_path / "t.pdf", ["first page text " * 5, "second page text " * 5])
        doc = pymupdf.open(pdf)
        doc.set_metadata({"title": "Quarterly Budget", "subject": "Finance review"})
        doc.saveIncr()
        doc.close()
        chunks = extractors.extract(pdf)
        assert chunks[0].text.startswith("Title: Quarterly Budget\nSubject: Finance review\n")
        assert "Title:" not in chunks[1].text

    def test_a_scanner_title_does_not_make_a_blank_scan_count(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        ocr.text = "' . , i l | ~"
        pdf = self.scans(tmp_path, 1)
        doc = pymupdf.open(pdf)
        doc.set_metadata({"title": "CamScanner 01-02-2026 10.30 document scan of the office"})
        doc.saveIncr()
        doc.close()
        chunks = build_with_ocr().extract(pdf)
        assert [c.low_content for c in chunks] == [True]
        assert "CamScanner 01-02-2026" in chunks[0].text  # still findable by its title

    def test_encrypted_pdf_is_rejected(self, extractors: ExtractorSet, tmp_path: Path) -> None:
        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), "secret")
        target = tmp_path / "enc.pdf"
        doc.save(target, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
        doc.close()
        with pytest.raises(ExtractError, match="encrypted"):
            extractors.extract(target)

    def test_corrupt_pdf_is_rejected(self, extractors: ExtractorSet, write: Writer) -> None:
        with pytest.raises(ExtractError, match="cannot open pdf"):
            extractors.extract(write("bad.pdf", b"%PDF-1.4 garbage"))


class TestDocx:
    def build(self, path: Path, with_image: bool = False) -> Path:
        import docx

        document = docx.Document()
        document.add_heading("Experience", level=1)
        document.add_paragraph("Built payment systems in Python.")
        document.add_heading("Skills", level=2)
        document.add_paragraph("FastAPI, PostgreSQL")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Language"
        table.cell(0, 1).text = "Years"
        table.cell(1, 0).text = "Python"
        table.cell(1, 1).text = "8"
        if with_image:
            document.add_picture(io.BytesIO(png_bytes()))
        document.save(path)
        return path

    def test_headings_become_symbols_and_tables_become_rows(
        self, extractors: ExtractorSet, tmp_path: Path
    ) -> None:
        chunks = extractors.extract(self.build(tmp_path / "r.docx"))
        by_symbol = {c.symbol: c.text for c in chunks}
        assert "Built payment systems" in by_symbol["Experience"]
        assert "FastAPI" in by_symbol["Experience > Skills"]
        assert "Language | Years" in by_symbol["table 1"]
        assert "Python | 8" in by_symbol["table 1"]

    def test_embedded_images_are_ocrd(self, build_with_ocr: Build, tmp_path: Path) -> None:
        chunks = build_with_ocr().extract(self.build(tmp_path / "img.docx", with_image=True))
        picture = next(c for c in chunks if c.kind == "image")
        assert "OAuth login flow" in picture.text

    def test_the_same_picture_twice_costs_one_ocr_and_one_allowance(
        self, build_with_ocr: Build, tmp_path: Path, ocr: FakeOcr
    ) -> None:
        import docx

        document = docx.Document()
        same = png_bytes()
        for _ in range(3):
            document.add_picture(io.BytesIO(same))
        document.add_picture(io.BytesIO(png_bytes((310, 310), "gray")))
        document.save(tmp_path / "dup.docx")
        chunks = build_with_ocr(ImageSettings(max_per_doc=2)).extract(tmp_path / "dup.docx")
        assert ocr.calls == 2  # one for the repeated picture, one for the different one
        assert len([c for c in chunks if c.kind == "image"]) == 2

    def test_long_section_is_split(
        self, with_chunking: Callable[..., ExtractorSet], tmp_path: Path
    ) -> None:
        import docx

        document = docx.Document()
        document.add_heading("Big", 1)
        for i in range(40):
            document.add_paragraph(f"paragraph number {i} " * 6)
        document.save(tmp_path / "big.docx")
        cfg = ChunkingSettings(max_chunk_chars=300, target_chunk_chars=150)
        chunks = with_chunking(cfg).extract(tmp_path / "big.docx")
        assert len([c for c in chunks if c.symbol == "Big"]) > 1

    def test_corrupt_docx_is_rejected(self, extractors: ExtractorSet, write: Writer) -> None:
        with pytest.raises(ExtractError, match="cannot open docx"):
            extractors.extract(write("bad.docx", b"not a zip"))

    def test_blank_docx_without_zip_media(self, extractors: ExtractorSet, tmp_path: Path) -> None:
        import docx

        docx.Document().save(tmp_path / "empty.docx")
        assert extractors.extract(tmp_path / "empty.docx") == []


class TestPptx:
    def build(self, path: Path, with_image: bool = False) -> Path:
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Quarterly roadmap"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text_frame.text = "Ship the indexer"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Q1"
        table.cell(0, 1).text = "Search"
        slide.notes_slide.notes_text_frame.text = "mention the GPU budget"
        if with_image:
            slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(4))
        deck.save(path)
        return path

    def test_slide_text_table_and_notes(self, extractors: ExtractorSet, tmp_path: Path) -> None:
        (chunk,) = extractors.extract(self.build(tmp_path / "d.pptx"))
        assert (chunk.symbol, chunk.page) == ("slide 1", 1)
        for expected in ("Quarterly roadmap", "Ship the indexer", "Q1 | Search", "Notes: mention"):
            assert expected in chunk.text

    def test_a_linked_picture_does_not_break_the_deck(
        self, extractors: ExtractorSet, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from pptx.shapes.picture import Picture

        def linked(_self: object) -> None:
            raise ValueError("no embedded image")  # python-pptx's answer for a linked picture

        monkeypatch.setattr(Picture, "image", property(linked))
        chunks = extractors.extract(self.build(tmp_path / "linked.pptx", with_image=True))
        assert "Ship the indexer" in chunks[0].text  # the slide's text is still indexed

    def test_pictures_are_ocrd(self, build_with_ocr: Build, tmp_path: Path) -> None:
        chunks = build_with_ocr().extract(self.build(tmp_path / "p.pptx", with_image=True))
        picture = next(c for c in chunks if c.kind == "image")
        assert picture.symbol == "image on slide 1"
        assert picture.page == 1

    def test_corrupt_pptx_is_rejected(self, extractors: ExtractorSet, write: Writer) -> None:
        with pytest.raises(ExtractError, match="cannot open pptx"):
            extractors.extract(write("bad.pptx", b"not a zip"))


class TestRepeatedLines:
    def test_lines_on_most_pages_are_furniture(self) -> None:
        pages = [
            ["Acme Corp", "Intro text", "Page 1 of 3"],
            ["Acme Corp", "Middle text", "Page 2 of 3"],
            ["Body only", "Page 3 of 3"],
        ]
        assert pdf_mod.repeated_lines(pages) == {"acme corp", "page # of #"}

    def test_two_pages_must_share_a_line(self) -> None:
        assert pdf_mod.repeated_lines([["only once"], ["something else"]]) == set()
        assert pdf_mod.repeated_lines([["Same  Header"], ["same header"]]) == {"same header"}

    def test_a_single_page_has_no_furniture(self) -> None:
        assert pdf_mod.repeated_lines([["Title", "Title"]]) == set()

    def test_a_line_twice_on_one_page_counts_once(self) -> None:
        pages = [["x", "x"], ["y"], ["z"], ["w"]]
        assert pdf_mod.repeated_lines(pages) == set()
