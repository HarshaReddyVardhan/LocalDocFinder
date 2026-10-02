"""DOCX: chunked by heading, tables as rows, embedded images through OCR."""

import logging
import zipfile
from collections.abc import Iterable
from pathlib import Path

from vector_embed.core.extractors.base import (
    KIND_DOC,
    Chunk,
    ExtractContext,
    ExtractError,
    Extractor,
    Untyped,
    register_extractor,
)
from vector_embed.core.extractors.chunking import split_by_lines
from vector_embed.core.extractors.image import image_chunk, image_key

logger = logging.getLogger(__name__)

_IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp", ".gif")


def _heading_level(style: str) -> int | None:
    lowered = style.lower()
    if not (lowered.startswith("heading") or style == "Title"):
        return None
    digits = "".join(ch for ch in style if ch.isdigit())
    return int(digits) if digits else 1


def _paragraph_chunks(ctx: ExtractContext, document: Untyped) -> list[Chunk]:
    cfg = ctx.chunking
    chunks: list[Chunk] = []
    stack: list[tuple[int, str]] = []
    path = ""
    buffer: list[str] = []

    def flush() -> None:
        text = "\n".join(buffer).strip()
        buffer.clear()
        if not text:
            return
        if len(text) <= cfg.max_chunk_chars:
            chunks.append(Chunk(text, KIND_DOC, path))
            return
        chunks.extend(
            Chunk(part, KIND_DOC, path)
            for part, _, _ in split_by_lines(
                text.splitlines(), 1, cfg.target_chunk_chars * 2, cfg.line_overlap
            )
        )

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        level = _heading_level(style)
        if level is not None:
            flush()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, text))
            path = " > ".join(t for _, t in stack)
        buffer.append(text)
    flush()
    return chunks


def _table_chunks(ctx: ExtractContext, document: Untyped) -> list[Chunk]:
    cfg = ctx.chunking
    chunks: list[Chunk] = []
    for index, table in enumerate(document.tables, 1):
        rows: list[str] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and (not cells or cells[-1] != text):  # merged cells repeat their text
                    cells.append(text)
            if cells:
                rows.append(" | ".join(cells))
        chunks.extend(
            Chunk(part, KIND_DOC, f"table {index}")
            for part, _, _ in split_by_lines(rows, 1, cfg.target_chunk_chars * 2, cfg.line_overlap)
        )
    return chunks


def _embedded_image_chunks(ctx: ExtractContext, path: Path) -> list[Chunk]:
    chunks: list[Chunk] = []
    budget = ctx.images.max_per_doc
    seen: set[str] = set()
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if budget <= 0:
                    break
                if name.startswith("word/media/") and name.lower().endswith(_IMAGE_SUFFIXES):
                    data = archive.read(name)
                    if image_key(data) in seen:
                        continue  # the same picture used twice: no second OCR, no allowance spent
                    budget -= 1
                    chunk = image_chunk(ctx, data, 0, seen, "embedded image")
                    if chunk:
                        chunks.append(chunk)
    except zipfile.BadZipFile:
        logger.debug("docx: not a zip archive: %s", path)
    return chunks


@register_extractor("docx")
class DocxExtractor(Extractor):
    name = "docx"
    priority = 30
    is_document = True

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def extract(self, path: Path) -> Iterable[Chunk]:
        import docx

        try:
            document = docx.Document(str(path))
        except Exception as exc:  # python-docx raises zip, XML and key errors for bad files
            raise ExtractError(f"cannot open docx: {exc}") from exc
        return [
            *_paragraph_chunks(self.ctx, document),
            *_table_chunks(self.ctx, document),
            *_embedded_image_chunks(self.ctx, path),
        ]
