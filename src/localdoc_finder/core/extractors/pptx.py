"""PPTX: one chunk per slide (text, tables, notes), pictures through OCR."""

import logging
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from localdoc_finder.core.extractors.base import (
    KIND_DOC,
    Chunk,
    ExtractError,
    Extractor,
    Untyped,
    register_extractor,
)
from localdoc_finder.core.extractors.image import image_chunk, image_key

logger = logging.getLogger(__name__)


def _walk(shapes: Untyped, group_type: Untyped) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == group_type:
            yield from _walk(shape.shapes, group_type)
        else:
            yield shape


def _embedded_blob(shape: Untyped) -> bytes | None:
    """A picture's bytes, or ``None`` for a linked picture (it names a file, not deck data)."""
    try:
        blob: bytes = shape.image.blob
    except (ValueError, KeyError, AttributeError):
        logger.debug("pptx: picture without embedded image data", exc_info=True)
        return None
    return blob


@register_extractor("pptx")
class PptxExtractor(Extractor):
    name = "pptx"
    priority = 30
    is_document = True

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    def extract(self, path: Path) -> Iterable[Chunk]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        try:
            presentation = Presentation(str(path))
        except Exception as exc:  # python-pptx raises zip, XML and key errors for bad files
            raise ExtractError(f"cannot open pptx: {exc}") from exc

        ctx = self.ctx
        chunks: list[Chunk] = []
        budget = ctx.images.max_per_doc
        seen: set[str] = set()
        for number, slide in enumerate(presentation.slides, 1):
            texts: list[str] = []
            for shape in _walk(slide.shapes, MSO_SHAPE_TYPE.GROUP):
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.strip())
                elif getattr(shape, "has_table", False) and shape.has_table:
                    texts.extend(
                        " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        for row in shape.table.rows
                    )
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and budget > 0:
                    blob = _embedded_blob(shape)
                    if blob is None or image_key(blob) in seen:
                        continue  # linked or missing picture, or one already read: free
                    budget -= 1
                    picture = image_chunk(ctx, blob, number, seen, f"image on slide {number}")
                    if picture:
                        chunks.append(picture)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    texts.append("Notes: " + notes)
            if texts:
                body = "\n".join(texts)[: ctx.chunking.max_chunk_chars]
                chunks.append(Chunk(body, KIND_DOC, f"slide {number}", page=number))
        return chunks
